from django.shortcuts import render, redirect, get_object_or_404
from django.template import loader
from django.http import HttpResponse
from django.views import View
import csv
from io import TextIOWrapper
from django.db import IntegrityError, transaction
from .models import Employee, Department, Role, Leave, Attendance, Performance, Payroll
from .forms import LeaveForm, EmployeeForm, BulkEmployeeUploadForm
from .forms import LeaveForm, EmployeeForm
from django.utils import timezone
from .forms import AttendanceForm, PerformanceForm, AttendanceUploadForm
from django.contrib import messages
from django.db.models import Q
from django.contrib.messages.views import SuccessMessageMixin
from django.urls import reverse_lazy
from django.core.paginator import Paginator
from django.views.generic.edit import CreateView, UpdateView, DeleteView
from django.views.generic.detail import DetailView
from authenication.models import CustomUser
import csv
import io
from authenication.models import CustomUser
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
import openpyxl
import pandas as pd
from django.contrib.auth.models import User
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from .forms import EmployeeProfileForm
from django.views import View
from django.shortcuts import render
from django.http import HttpResponse
from hr.models import Payroll
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from decimal import Decimal
from datetime import datetime
from io import BytesIO

class DashboardView(View):
    def get(self, request):
        template = loader.get_template('hr/pages/dashboard.html')
        employees = Employee.objects.all()
        context = {
            'employee_count': employees.count(),
            'department_count': Department.objects.count(),
        }
        return HttpResponse(template.render(context, request))


class EmployeesView(View):
    def get(self, request):
        template = loader.get_template('hr/pages/employees.html')
        employees = Employee.objects.all()
        context = {
            'employees': employees,
            'department_count': Department.objects.count(),
        }
        return HttpResponse(template.render(context, request))

class EmployeeCreateView(SuccessMessageMixin, CreateView):
    model = Employee
    form_class = EmployeeForm
    template_name = 'hr/pages/employee_form.html'
    success_url = reverse_lazy('employees')
    success_message = 'Employee created successfully!'

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        form.fields['user'].queryset = CustomUser.objects.filter(employee__isnull=True)
        return form


class EmployeeUpdateView(SuccessMessageMixin, UpdateView):
    model = Employee
    form_class = EmployeeForm
    template_name = 'hr/pages/employee_form.html'
    success_url = reverse_lazy('employees')
    success_message = 'Employee updated successfully!'

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        form.fields['user'].queryset = CustomUser.objects.filter(
            Q(employee__isnull=True) | Q(employee=self.object)
        )
        return form


class EmployeeDeleteView(SuccessMessageMixin, DeleteView):
    model = Employee
    template_name = 'hr/pages/employee_confirm_delete.html'
    success_url = reverse_lazy('employees')
    success_message = 'Employee deleted successfully!'

    def delete(self, request, *args, **kwargs):
        response = super().delete(request, *args, **kwargs)
        messages.success(request, self.success_message)
        return response


class EmployeeDetailView(DetailView):
    model = Employee
    template_name = 'hr/pages/employee_detail.html'
    context_object_name = 'employee'

#sethu
class LeaveListView(View):
    def get(self, request):
        template = loader.get_template('hr/pages/leave.html')
        leaves = Leave.objects.all()
        context = {
            'leaves': leaves
        }
        return HttpResponse(template.render(context, request))

def leave_list(request):

    leaves = Leave.objects.all()

    context = {
        "leaves": leaves
    }

    return render(request, "hr/pages/leave.html", context)


def leave_create(request):
    if request.method == "POST":
        form = LeaveForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('leave')
    else:
        form = LeaveForm()

    return render(request, "hr/pages/create.html", {"form": form})


def leave_view(request, id):
    leave = get_object_or_404(Leave, id=id)
    return render(request, "hr/pages/leave_view.html", {"leave": leave})


def leave_update(request, id):
    leave = get_object_or_404(Leave, id=id)

    if request.method == "POST":
        form = LeaveForm(request.POST, instance=leave)
        if form.is_valid():
            form.save()
            return redirect('leave')
    else:
        form = LeaveForm(instance=leave)

    return render(request, "hr/pages/leave_update.html", {"form": form})


def leave_delete(request, id):
    leave = get_object_or_404(Leave, id=id)

    if request.method == "POST":
        leave.delete()
        return redirect('leave')

    return redirect('leave')
#bulk_leave_approval
@login_required
def bulk_update_leave_status(request):

    if request.method == "POST":

        leave_ids = request.POST.getlist("leave_ids")
        action = request.POST.get("action")

        if not leave_ids:
            messages.error(request, "No leave requests selected")
            return redirect("leave-list")

        Leave.objects.filter(id__in=leave_ids).update(status=action)

        messages.success(
            request,
            f"{len(leave_ids)} leave requests updated to {action}"
        )

        return redirect("leave-list")


@login_required
def bulk_leave_list(request):

    leaves = Leave.objects.filter(status="Pending")

    return render(
        request,
        "hr/pages/leave_list.html",
        {
            "leaves": leaves
        }
    )

#bulk_leave_upload
@login_required
def bulk_leave_upload(request):
    if request.method == "POST":
        file = request.FILES["file"]
        decoded_file = file.read().decode("utf-8")
        io_string = io.StringIO(decoded_file)
        reader = csv.DictReader(io_string)

        errors = []

        for idx, row in enumerate(reader, start=2):
            employee_name = row.get("employee_name", "").strip()
            if not employee_name:
                errors.append(f"Line {idx}: Employee name is missing.")
                continue

            # Split first and last name
            try:
                first_name, last_name = employee_name.split(" ", 1)
                employee = Employee.objects.filter(first_name=first_name, last_name=last_name).first()
            except ValueError:
                employee = None

            if employee:
                # ✅ Pass employee object to employee field, not first_name/last_name
                Leave.objects.create(
                    employee=employee,
                    start_date=row.get("start_date"),
                    end_date=row.get("end_date"),
                    reason=row.get("reason"),
                    status=row.get("status")
                )
            else:
                errors.append(f"Line {idx}: Employee '{employee_name}' not found.")

        if errors:
            return render(request, "hr/pages/bulk_leave_upload.html", {"errors": errors})
        else:
            messages.success(request, "CSV file uploaded successfully!")
            return redirect("leave")

    return render(request, "hr/pages/bulk_leave_upload.html")


@login_required
def download_leave_csv_template(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="leave_template.csv"'

    writer = csv.writer(response)
  
    writer.writerow(['first_name', 'start_date', 'end_date', 'reason', 'status'])

    first_employee = Employee.objects.first()
    if first_employee:
        writer.writerow([first_employee.name, '2026-03-10', '2026-03-12', 'Sick', 'Pending'])

    return response

@login_required
def view_export_leave(request):
    leave = Leave.objects.all().select_related('employee')
    return render(request, "hr/pages/view_export_leave.html", {"leave": leave})

# Export Excel
@login_required
def export_leave_excel(request):
    leave = Leave.objects.all().select_related('employee')
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Leave"

    # Header
    headers = ['Employee', 'Start Date', 'End Date', 'Reason', 'Status']
    sheet.append(headers)

    # Data
    for leave in leave:
        sheet.append([
            f"{leave.employee.first_name} {leave.employee.last_name}",
            leave.start_date,
            leave.end_date,
            leave.reason,
            leave.status
        ])

    # Prepare response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=leave_list.xlsx'
    workbook.save(response)
    return response

# Export PDF
@login_required
def export_leave_pdf(request):
    leave = Leave.objects.all().select_related('employee')
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="leave_list.pdf"'

    buffer = BytesIO()
    p = canvas.Canvas(buffer)

    y = 800
    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, "Leave List")
    y -= 30

    # Table header
    p.setFont("Helvetica-Bold", 10)
    headers = ['Employee', 'Start Date', 'End Date', 'Reason', 'Status']
    x_positions = [50, 200, 300, 400, 500]
    for i, header in enumerate(headers):
        p.drawString(x_positions[i], y, header)
    y -= 20

    p.setFont("Helvetica", 10)
    for leave in leave:
        if y < 50:
            p.showPage()
            y = 800
        p.drawString(x_positions[0], y, f"{leave.employee.first_name} {leave.employee.last_name}")
        p.drawString(x_positions[1], y, str(leave.start_date))
        p.drawString(x_positions[2], y, str(leave.end_date))
        p.drawString(x_positions[3], y, leave.reason)
        p.drawString(x_positions[4], y, leave.status)
        y -= 20

    p.save()
    pdf = buffer.getvalue()
    buffer.close()
    response.write(pdf)
    return response

#Mbali's code

class Attendance_View(View):
    def get(self, request):
        template = loader.get_template('hr/pages/attendance.html')
        
        today = timezone.now().date()
        recent_attendances = Attendance.objects.select_related(
            'employee__department'  
        ).prefetch_related(
            'employee'  
        ).filter(
            employee__isnull=False 
        ).order_by('-date', '-created_at')[:50]
        total_employees = Employee.objects.count()
        present_today = Attendance.objects.filter(date=today, status='present').count()
        late_today = Attendance.objects.filter(date=today, status='late').count()
        absent_today = Attendance.objects.filter(date=today, status='absent').count()
        
        stats = {
            'present': present_today,
            'late': late_today,
            'absent': absent_today,
            'total': total_employees
        }
        
        context = {
            'recent_attendances': recent_attendances,
            'stats': stats,
            'today': today,
        }
        return HttpResponse(template.render(context, request))

def attendance_create(request):
    if request.method == 'POST':
        form = AttendanceForm(request.POST)
        if form.is_valid():
            attendance = form.save()
            print("SAVED:", attendance)
            return redirect('attendance')
        else:
            print(form.errors)
    else:
        form = AttendanceForm()

    return render(request, 'hr/pages/attendance_form.html', {
        'form': form,
        'title': 'Add Attendance'
    })

def attendance_edit(request, pk):
    attendance = get_object_or_404(Attendance, pk=pk)

    if request.method == 'POST':
        form = AttendanceForm(request.POST, instance=attendance)
        if form.is_valid():
            form.save()
            return redirect('attendance')
    else:
        form = AttendanceForm(instance=attendance)

    return render(request, 'hr/pages/attendance_form.html', {
        'form': form,
        'title': 'Edit Attendance'
    })

def attendance_delete(request, pk):
    attendance = get_object_or_404(Attendance, pk=pk)

    if request.method == 'POST':
        attendance.delete()
        return redirect('attendance')

    return render(request, 'hr/pages/attendance_confirm_delete.html', {
        'attendance': attendance
    })

def attendance_list(request):
    recent_attendances = Attendance.objects.select_related('employee').order_by('-date', '-id')

    return render(request, 'hr/pages/attendance.html', {
        'recent_attendances': recent_attendances
    })


def bulk_upload_attendance(request):
    if request.method == 'POST':
        form = AttendanceUploadForm(request.POST, request.FILES)

        if form.is_valid():
            csv_file = request.FILES['file']

            if not csv_file.name.endswith('.csv'):
                messages.error(request, "Please upload a CSV file.")
                return redirect('attendance')

            try:
                data = csv_file.read().decode('utf-8')
            except UnicodeDecodeError:
                csv_file.seek(0)
                data = csv_file.read().decode('ISO-8859-1')

            io_string = io.StringIO(data)
            reader = csv.DictReader(io_string)

            today = timezone.localdate()

            uploaded_employees = []

            for row in reader:
                first_name = row.get('first_name')
                last_name = row.get('last_name')

                if not first_name or not last_name:
                    continue

                try:
                    employee = Employee.objects.get(
                        first_name=first_name.strip(),
                        last_name=last_name.strip()
                    )

                    Attendance.objects.update_or_create(
                        employee=employee,
                        date=today,
                        defaults={'status': 'Present'}
                    )

                    uploaded_employees.append(employee.id)

                except Employee.DoesNotExist:
                    continue

            all_employees = Employee.objects.exclude(id__in=uploaded_employees)

            for employee in all_employees:
                Attendance.objects.get_or_create(
                    employee=employee,
                    date=today,
                    defaults={'status': 'Absent'}
                )

            messages.success(
                request,
                f"Attendance uploaded successfully! ({len(uploaded_employees)} Present, {all_employees.count()} Absent)"
            )

            return redirect('attendance')

    else:
        form = AttendanceUploadForm()

    return render(request, 'hr/pages/bulk_upload.html', {'form': form})

def attendance_bulk_delete(request):
    if request.method == "POST":
        ids = request.POST.getlist('selected_attendances')
        if ids:
            Attendance.objects.filter(id__in=ids).delete()
            messages.success(request, f"{len(ids)} attendance records deleted successfully!")
        else:
            messages.error(request, "No records selected.")
    return redirect('attendance')

def attendance_report(request):
    attendances = Attendance.objects.select_related('employee').all().order_by('-date')

    return render(request, 'hr/pages/attendance_report.html', {
        'attendances': attendances
    })

def export_attendance_excel(request):
    attendances = Attendance.objects.select_related('employee').all()

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Attendance Report"

    sheet.append(["Employee", "Date", "Status", "Check In", "Check Out"])

    for attendance in attendances:
        sheet.append([
            f"{attendance.employee.first_name} {attendance.employee.last_name}",
            str(attendance.date),
            attendance.status,
            str(attendance.check_in_time),
            str(attendance.check_out_time)
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=attendance_report.xlsx'

    workbook.save(response)

    return response

def export_attendance_pdf(request):

    attendances = Attendance.objects.select_related('employee').all()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="attendance_report.pdf"'

    p = canvas.Canvas(response)

    y = 800
    p.drawString(200, y, "Attendance Report")

    y -= 40

    for attendance in attendances:
        text = f"{attendance.employee.first_name} {attendance.employee.last_name} | {attendance.date} | {attendance.status}"
        p.drawString(50, y, text)
        y -= 20

    p.showPage()
    p.save()

    return response

#Lusanda code will go here

def departments(request):
    context = {
        'departments': Department.objects.all(),
        'department_count': Department.objects.count(),
    }
    return render(request, 'hr/pages/departments.html', context)


def performance_list(request):
    search_query = request.GET.get('search', '')

    performance_qs = Performance.objects.select_related(
        'employee',
        'department'
    ).order_by('-rating', '-id')

    if search_query:
        performance_qs = performance_qs.filter(
            Q(employee__first_name__icontains=search_query) |
            Q(employee__last_name__icontains=search_query)
        )

    paginator = Paginator(performance_qs, 10)
    page_number = request.GET.get('page')
    performances = paginator.get_page(page_number)

    return render(request, 'hr/pages/performance.html', {
        'performances': performances,
        'search_query': search_query
    })


def add_performance(request):
    if request.method == 'POST':
        form = PerformanceForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Performance record added successfully.')
            return redirect('performance')
    else:
        form = PerformanceForm()

    return render(request, 'hr/pages/performance_add.html', {
        'form': form,
        'title': 'Add Performance'
    })


def edit_performance(request, pk):
    performance = get_object_or_404(Performance, pk=pk)

    if request.method == 'POST':
        form = PerformanceForm(request.POST, instance=performance)

        if form.is_valid():
            form.save()  # Just save normally

            messages.success(request, "Performance review updated successfully.")
            return redirect('performance')
    else:
        form = PerformanceForm(instance=performance)

    return render(request, 'hr/pages/performance_edit.html', {
        'form': form,
        'title': 'Edit Performance'
    })

def delete_performance(request, pk):
    performance = get_object_or_404(Performance, pk=pk)

    if request.method == 'POST':
        performance.delete()
        messages.success(request, 'Performance record deleted successfully.')
        return redirect('performance')
    return render(request, 'hr/pages/performance_confirm_delete.html', {
        'performance': performance
    })

#Senzo's Code 
class DepartmentsView(View):
    def get(self, request):
        departments = Department.objects.prefetch_related('employee_set').all()
        context = {
            'departments': departments,
        }
        return render(request, 'hr/pages/departments.html', context)

class DepartmentCreateView(SuccessMessageMixin, CreateView):
    model = Department
    fields = ['name', 'description']
    template_name = 'hr/pages/department_form.html'
    success_url = reverse_lazy('department_list')
    success_message = 'Department created successfully!'

class DepartmentUpdateView(SuccessMessageMixin, UpdateView):
    model = Department
    fields = ['name', 'description']
    template_name = 'hr/pages/department_form.html'
    success_url = reverse_lazy('department_list')
    success_message = 'Department updated successfully!'

class DepartmentDeleteView(SuccessMessageMixin, DeleteView):
    model = Department
    template_name = 'hr/pages/department_confirm_delete.html'
    success_url = reverse_lazy('department_list')
    success_message = 'Department deleted successfully!'

class DepartmentBulkCreateView(SuccessMessageMixin, View):
    template_name = 'hr/pages/department_bulk_form.html'
    
    def get(self, request):
        return render(request, self.template_name)
    
    def post(self, request):
        departments_data = []
        
        i = 1
        while True:
            name_field = f'dept_name_{i}'
            desc_field = f'dept_desc_{i}'
            
            name = request.POST.get(name_field, '').strip()
            description = request.POST.get(desc_field, '').strip()
            
            if not name:
                break
                
            departments_data.append({
                'name': name,
                'description': description
            })
            i += 1
        
        if not departments_data:
            messages.error(request, 'Please add at least one department.')
            return render(request, self.template_name)
        
        existing_names = set(Department.objects.values_list('name', flat=True))
        new_depts_data = [d for d in departments_data if d['name'] not in existing_names]
        
        if not new_depts_data:
            messages.warning(request, 'All department names already exist.')
            return redirect('department_list')
        
        try:
            departments = [Department(name=d['name'], description=d['description']) 
                          for d in new_depts_data]
            Department.objects.bulk_create(departments)
            messages.success(request, f'Created {len(new_depts_data)} department(s)!')
        except Exception as e:
            messages.error(request, f'Error: {str(e)}')
        
        return redirect('department_list')

class DepartmentPDFExportView(View):
    def get(self, request):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        
        styles = getSampleStyleSheet()
        elements.append(Paragraph("All Departments Report", styles['Title']))
        elements.append(Spacer(1, 0.5 * inch))
        
        departments = Department.objects.all().prefetch_related('employee_set')
        data = [['Department Name', 'Description', 'Total Employees']]
        
        for dept in departments:
            data.append([
                dept.name,
                dept.description or "No description", 
                str(dept.employee_set.count())
            ])
        
        table = Table(data, colWidths=[2.8*inch, 3.5*inch, 1.5*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'), 
            ('ALIGN', (2, 0), (2, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
            ('RIGHTPADDING', (0, 0), (-1, -1), 12),
        ]))
        elements.append(table)
        
        elements.append(Spacer(1, 0.3 * inch))
        elements.append(Paragraph(
            f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", 
            styles['Normal']
        ))
        
        doc.build(elements)
        buffer.seek(0)
        
        return FileResponse(
            buffer, 
            as_attachment=True, 
            filename=f"departments_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
            content_type='application/pdf'
        )

#Zee's code will go here
class PayrollView(View):
    def get(self, request):
        payrolls = Payroll.objects.select_related('employee').all()
        employees = Employee.objects.all()
        context = {
            'payrolls': payrolls,
            'employees': employees,
            'total_employees': employees.count(),
            'total_payroll': payrolls.count(),
        }
        return render(request, 'hr/pages/payroll.html', context)


class AddPayrollView(View):
    def post(self, request):
        employee_id = request.POST.get('employee')
        month = request.POST.get('month')

        try:
            employee = Employee.objects.get(id=employee_id)
        except (Employee.DoesNotExist, ValueError, TypeError):
            messages.error(request, "Invalid employee selected.")
            return redirect('payroll')

        try:
            basic_salary = float(request.POST.get('basic_salary') or 0)
            bonus = float(request.POST.get('bonus') or 0)
            deductions = float(request.POST.get('deductions') or 0)
        except ValueError:
            messages.error(request, "Invalid salary, bonus, or deductions value.")
            return redirect('payroll')

        if Payroll.objects.filter(employee=employee, month=month).exists():
            messages.error(request, f"Payroll for {employee.first_name} {employee.last_name} in {month} already exists!")
            return redirect('payroll')

        Payroll.objects.create(
            employee=employee,
            basic_salary=basic_salary,
            bonus=bonus,
            deductions=deductions,
            month=month
        )
        messages.success(request, "Payroll added successfully.")
        return redirect('payroll')

class EditPayrollView(View):
    def get(self, request, pk):
        payroll = get_object_or_404(Payroll, pk=pk)
        employees = Employee.objects.all()
        context = {
            'payroll': payroll,
            'employees': employees,
        }
        return render(request, 'hr/pages/edit_payroll.html', context)

    def post(self, request, pk):
        payroll = get_object_or_404(Payroll, pk=pk)
        payroll.employee_id = request.POST.get('employee')
        payroll.basic_salary = float(request.POST.get('basic_salary'))
        payroll.bonus = float(request.POST.get('bonus') or 0)
        payroll.deductions = float(request.POST.get('deductions') or 0)
        payroll.month = request.POST.get('month')
        payroll.save()
        messages.success(request, "Payroll updated successfully.")
        return redirect('payroll')


class DeletePayrollView(View):
    def get(self, request, pk):
        payroll = get_object_or_404(Payroll, pk=pk)
        return render(request, 'hr/pages/delete_payroll.html', {'payroll': payroll})

    def post(self, request, pk):
        payroll = get_object_or_404(Payroll, pk=pk)
        payroll.delete()
        messages.success(request, "Payroll deleted successfully.")
        return redirect('payroll')


class ViewPayrollReport(View):
    template_name = 'hr/pages/view_payroll.html'

    def get(self, request):
        payrolls = Payroll.objects.select_related('employee').all()
        for p in payrolls:
            p.total_salary = (p.basic_salary or 0) + (p.bonus or 0) - (p.deductions or 0)
        return render(request, self.template_name, {'payrolls': payrolls})

    def post(self, request):
        payrolls = Payroll.objects.select_related('employee').all()
        for p in payrolls:
            p.total_salary = (p.basic_salary or 0) + (p.bonus or 0) - (p.deductions or 0)

        if 'export_pdf' in request.POST:
            return self.export_pdf(payrolls)
        if 'export_excel' in request.POST:
            return self.export_excel(payrolls)
        if 'export_word' in request.POST:
            return self.export_word(payrolls)

        return render(request, self.template_name, {'payrolls': payrolls})

    def export_pdf(self, payrolls):
        from io import BytesIO
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []

        styles = getSampleStyleSheet()
        elements.append(Paragraph("Monthly Payroll Report", styles['Title']))
        elements.append(Spacer(1, 20))

        data = [["Employee", "Month", "Basic Salary", "Bonus", "Deductions", "Total Salary"]]
        for p in payrolls:
            data.append([
                f"{p.employee.first_name} {p.employee.last_name}",
                p.month,
                f"{p.basic_salary:.2f}",
                f"{p.bonus:.2f}",
                f"{p.deductions:.2f}",
                f"{p.total_salary:.2f}"
            ])

        table = Table(data, hAlign='LEFT', colWidths=[120, 80, 80, 60, 80, 80])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#34495e')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 12),
            ('BOTTOMPADDING', (0,0), (-1,0), 10),
            ('BACKGROUND', (0,1), (-1,-1), colors.whitesmoke),
            ('GRID', (0,0), (-1,-1), 0.5, colors.black),
        ]))
        elements.append(table)

        def add_watermark(canvas_obj, doc_obj):
            canvas_obj.saveState()
            canvas_obj.setFont('Helvetica-Bold', 40)
            canvas_obj.setFillColorRGB(0.6, 0.3, 0.8, alpha=0.3)
            canvas_obj.drawString(450, 800, "HRMS")
            canvas_obj.restoreState()

        doc.build(elements, onFirstPage=add_watermark, onLaterPages=add_watermark)
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="monthly_payroll_report.pdf"'
        return response

    def export_excel(self, payrolls):
        import pandas as pd
        from io import BytesIO

        data = []
        for p in payrolls:
            data.append({
                'Employee': f"{p.employee.first_name} {p.employee.last_name}",
                'Month': p.month,
                'Basic Salary': p.basic_salary,
                'Bonus': p.bonus,
                'Deductions': p.deductions,
                'Total Salary': p.total_salary
            })

        df = pd.DataFrame(data)
        buffer = BytesIO()
        with pd.ExcelWriter(buffer, engine='xlsxwriter') as writer:
            df.to_excel(writer, index=False, sheet_name='Payroll')
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="monthly_payroll_report.xlsx"'
        return response

         # ---------------- Word Export ----------------
    def export_word(self, payrolls):
        document = Document()
        document.add_heading('Monthly Payroll Report', level=1)

        table = document.add_table(rows=1, cols=6)
        headers = ["Employee", "Month", "Basic Salary", "Bonus", "Deductions", "Total Salary"]

        # Header row
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            self.set_cell_background(hdr_cells[i], "FF0000")  # red header
            for paragraph in hdr_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                    run.font.size = Pt(11)

        # Data rows with alternate shading
        for idx, p in enumerate(payrolls):
            row_cells = table.add_row().cells
            row_cells[0].text = f"{p.employee.first_name} {p.employee.last_name}"
            row_cells[1].text = str(p.month)
            row_cells[2].text = f"R{p.basic_salary:.2f}"
            row_cells[3].text = f"R{p.bonus:.2f}"
            row_cells[4].text = f"R{p.deductions:.2f}"
            row_cells[5].text = f"R{p.total_salary:.2f}"

            if idx % 2 == 0:  # light gray for alternate rows
                for cell in row_cells:
                    self.set_cell_background(cell, "F5F5F5")

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="monthly_payroll_report.docx"'
        document.save(response)
        return response

    # ---------------- Helper to set Word cell background ----------------
    def set_cell_background(self, cell, fill):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), fill)
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        
def bulk_upload_payroll(request):
    if request.method != "POST":
        return redirect('payroll')

    file = request.FILES.get("file")
    if not file:
        messages.error(request, "Please upload an Excel file.")
        return redirect('payroll')

    try:
        df = pd.read_excel(file)
    except Exception as e:
        messages.error(request, f"Failed to read Excel file: {str(e)}")
        return redirect('payroll')

    df.columns = df.columns.str.strip().str.lower()
    required_columns = ["employee_name", "basic_salary", "bonus", "deductions", "month"]

    for col in required_columns:
        if col not in df.columns:
            messages.error(request, f"Missing column: {col}")
            return redirect('payroll')

    success_count = 0
    for index, row in df.iterrows():
        employee_name = str(row.get("employee_name", "")).strip()
        if not employee_name:
            continue

        try:
            first, *last = employee_name.split()
            last = " ".join(last)

            employee = Employee.objects.filter(
                first_name__iexact=first,
                last_name__iexact=last
            ).first()

            if not employee:
                messages.error(request, f"Row {index+2}: Employee '{employee_name}' not found.")
                continue

            basic_salary = float(row.get("basic_salary") or 0)
            bonus = float(row.get("bonus") or 0)
            deductions = float(row.get("deductions") or 0)
            month = str(row.get("month") or "").strip()

            if Payroll.objects.filter(employee=employee, month=month).exists():
                messages.error(request, f"Row {index+2}: Payroll already exists for {employee_name} ({month}).")
                continue

            Payroll.objects.create(
                employee=employee,
                basic_salary=basic_salary,
                bonus=bonus,
                deductions=deductions,
                month=month
            )
            success_count += 1

        except Exception as e:
            messages.error(request, f"Row {index+2}: Invalid data ({str(e)})")

    if success_count:
        messages.success(request, f"{success_count} payroll record(s) uploaded successfully.")
    return redirect('payroll')


User = get_user_model()

def bulk_upload_performance(request):
    if request.method == "POST":
        csv_file = request.FILES['file']
        if not csv_file.name.endswith('.csv'):
            messages.error(request, "File must be CSV")
            return redirect('bulk_upload_performance')

        try:
            decoded_file = csv_file.read().decode('utf-8-sig').splitlines()
            delimiter = ';' if ';' in decoded_file[0] else ','
            reader = csv.DictReader(decoded_file, delimiter=delimiter)

            print("CSV headers found:", reader.fieldnames)
            for row in reader:
                try:
                    department, _ = Department.objects.get_or_create(
                        name=row['department'].strip()
                    )

                    role_obj, _ = Role.objects.get_or_create(
                        name=row['role'].strip()
                    )

                    username = f"{row['employee_name'].strip().lower()}_{row['employee_surname'].strip().lower()}"
                    user_email = f"{username}@example.com"
                    user, _ = User.objects.get_or_create(
                        username=username,
                        defaults={'email': user_email}
                    )

                    employee_email = f"{username}@example.com"
                    employee, _ = Employee.objects.get_or_create(
                        user=user,
                        first_name=row['employee_name'].strip(),
                        last_name=row['employee_surname'].strip(),
                        defaults={
                            'department': department,
                            'role': role_obj,
                            'email': employee_email
                        }
                    )

                    Performance.objects.create(
                        employee=employee,
                        department=department,
                        role=role_obj,
                        rating=int(row['rating']),
                        remarks=row['remarks'].strip()
                    )

                except KeyError as ke:
                    print(f"Missing column in CSV: {ke}")
                except Exception as e:
                    print("Error:", e)

            messages.success(request, "Performance records uploaded successfully")
        except Exception as e:
            messages.error(request, f"Error reading CSV: {e}")
        return redirect('performance')
    return render(request, 'hr/pages/performance_bulk_upload.html')

def export_performance(request, export_format):
    performances = Performance.objects.select_related('employee', 'department').all()

    if export_format == 'excel':
        # Excel export
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=performance.xlsx'

        wb = Workbook()
        ws = wb.active
        ws.title = "Performance"

        ws.append(['Employee', 'Department', 'Role', 'Rating', 'Remarks'])

        for perf in performances:
            ws.append([
                f"{perf.employee.first_name} {perf.employee.last_name}",
                perf.department.name,
                perf.role,
                str(perf.rating),
                perf.remarks
            ])

        wb.save(response)
        return response

    elif export_format == 'pdf':
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="performance.pdf"'

        p = canvas.Canvas(response, pagesize=letter)
        width, height = letter
        y = height - 50

        p.setFont("Helvetica-Bold", 14)
        p.drawString(50, y, "Performance Records")
        y -= 30
        p.setFont("Helvetica", 12)

        headers = ['Employee', 'Department', 'Role', 'Rating', 'Remarks']
        p.drawString(50, y, ' | '.join(headers))
        y -= 20

        for perf in performances:
            row = [
                f"{perf.employee.first_name} {perf.employee.last_name}",
                perf.department.name,
                perf.role,
                str(perf.rating),
                perf.remarks
            ]
            p.drawString(50, y, ' | '.join(row))
            y -= 20
            if y < 50:
                p.showPage()
                y = height - 50

        p.showPage()
        p.save()
        return response

    elif export_format == 'word':
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="performance.docx"'

        doc = Document()
        doc.add_heading('Performance Records', 0)

        table = doc.add_table(rows=1, cols=5)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Employee'
        hdr_cells[1].text = 'Department'
        hdr_cells[2].text = 'Role'
        hdr_cells[3].text = 'Rating'
        hdr_cells[4].text = 'Remarks'

        for perf in performances:
            row_cells = table.add_row().cells
            row_cells[0].text = f"{perf.employee.first_name} {perf.employee.last_name}"
            row_cells[1].text = perf.department.name
            row_cells[2].text = perf.role
            row_cells[3].text = str(perf.rating)
            row_cells[4].text = perf.remarks

        doc.save(response)
        return response

    else:
        messages.error(request, "Invalid export format")
        return redirect('performance')








class EmployeeProfilePicureView(View):
    def get(self, request, pk):
        employee = get_object_or_404(Employee, pk=pk)
        context = {
            'form': EmployeeProfileForm(instance=employee),
        }
        return render(request, 'hr/pages/profile.html', context)

    def post(self, request, pk):
        employee = get_object_or_404(Employee, pk=pk)

        form = EmployeeProfileForm(
            request.POST,
            request.FILES,
            instance=employee
        )

        if form.is_valid():
            # save only if a new image was uploaded
            if request.FILES.get('profile_picture'):
                employee.profile_picture = request.FILES['profile_picture']
                employee.save()

            return redirect('employee_profile', pk=employee.pk)

        context = {
            'form': form,
            'error': form.errors,
        }
        return render(request, 'hr/pages/profile.html', context)


class EmployeeBulkUploadView(View):
    template_name = 'hr/pages/employee_bulk_upload.html'

    def get(self, request):
        form = BulkEmployeeUploadForm()
        return render(request, self.template_name, {'form': form})

    def post(self, request):
        form = BulkEmployeeUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, self.template_name, {'form': form})

        csv_file = form.cleaned_data['file']
        reader = csv.DictReader(TextIOWrapper(csv_file.file, encoding='utf-8-sig'))

        required_columns = {'username', 'email', 'first_name', 'last_name', 'department', 'role'}
        header_cols = set([c.strip() for c in (reader.fieldnames or [])])
        missing_cols = required_columns - header_cols

        if missing_cols:
            form.add_error('file', f"Missing columns: {', '.join(sorted(missing_cols))}")
            return render(request, self.template_name, {'form': form})

        results = {'created': 0, 'failed': 0, 'errors': []}

        for row_num, row in enumerate(reader, start=2):
            username = (row.get('username') or '').strip()
            email = (row.get('email') or '').strip()
            first_name = (row.get('first_name') or '').strip()
            last_name = (row.get('last_name') or '').strip()
            dept_name = (row.get('department') or '').strip()
            role_name = (row.get('role') or '').strip()
            password = (row.get('password') or '').strip()

            if not all([username, email, first_name, last_name, dept_name, role_name]):
                results['failed'] += 1
                results['errors'].append(f"Line {row_num}: Missing required fields.")
                continue

            department = Department.objects.filter(name__iexact=dept_name).first()
            if not department:
                department = Department.objects.create(name=dept_name, description="")

            role = Role.objects.filter(name__iexact=role_name).first()
            if not role:
                role = Role.objects.create(name=role_name, description="")

            try:
                with transaction.atomic():
                    user = CustomUser(
                        username=username,
                        email=email,
                        role='employee',
                        first_name=first_name,
                        last_name=last_name,
                    )
                    if password:
                        user.set_password(password)
                    else:
                        user.set_unusable_password()
                    user.save()

                    Employee.objects.create(
                        user=user,
                        first_name=first_name,
                        last_name=last_name,
                        email=email,
                        department=department,
                        role=role,
                    )

                results['created'] += 1

            except IntegrityError:
                results['failed'] += 1
                results['errors'].append(f"Line {row_num}: Duplicate username or email.")
            except Exception as e:
                results['failed'] += 1
                results['errors'].append(f"Line {row_num}: {e}")

        if results['created']:
            messages.success(request, f"Created {results['created']} employees.")
        if results['failed']:
            messages.error(request, f"Failed {results['failed']} rows.")

        return render(request, self.template_name, {
            'form': BulkEmployeeUploadForm(),
            'results': results
        })

class export_employee_profileView(View):
    def get(self, request):
        employees = Employee.objects.select_related('user', 'department', 'role').all()
        if request.GET.get('format') == 'pdf':
            response = HttpResponse(content_type='application/pdf')
            response['Content-Disposition'] = 'attachment; filename="employee_profiles.pdf"'

            p = canvas.Canvas(response)
            y = 800
            p.setFont('Helvetica-Bold', 16)
            p.drawString(200, y, 'Employee Profiles')
            y -= 30
            p.setFont('Helvetica', 11)

            for employee in employees:
                full_name = f"{employee.first_name} {employee.last_name}".strip()
                email = employee.email or (employee.user.email if employee.user else "")
                department = employee.department.name if employee.department else ""
                role = employee.role.name if employee.role else ""

                lines = [
                    f"Name: {full_name}",
                    f"Email: {email}",
                    f"Department: {department}",
                    f"Role: {role}",
                    ""
                ]

                for line in lines:
                    if y <= 60:
                        p.showPage()
                        y = 800
                        p.setFont('Helvetica', 11)
                    p.drawString(50, y, line)
                    y -= 18

            p.save()
            return response

        context = {
            'employees': employees
        }
        return render(request, 'hr/pages/export_employee_profile.html', context)
